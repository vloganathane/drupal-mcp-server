"""Unit tests for the export tools added in Wave 3.

The contract under test: rows are paged per bundle over JSON:API (a bare
entity-type collection merges bundles behind one shared offset, which
would skip rows, so the tools always walk `type--bundle` endpoints),
destinations are confined to safe roots, and the response carries file
statistics rather than the rows themselves. Taxonomy usage SQL is
asserted by exact string: the field names come from config, but the
queries are still validated before they run.
"""

import asyncio
import csv as csv_module
import json
import tempfile
from pathlib import Path
from typing import Any

import drupal_mcp_server.tools.exports as exports


def _tool_functions(module: Any, **deps: Any) -> dict[str, Any]:
    class FakeMCP:
        def __init__(self) -> None:
            self.fns: dict[str, Any] = {}

        def tool(self, *args: Any, **kwargs: Any):
            def decorate(fn):
                self.fns[fn.__name__] = fn
                return fn

            return decorate

    mcp = FakeMCP()
    module.register_tools(mcp, **deps)
    return mcp.fns


class _Settings:
    def __init__(self, drupal_root: str | None) -> None:
        self.DRUPAL_ROOT = drupal_root


class _JSONAPI:
    """Paged JSON:API with include resolution from a uuid index."""

    def __init__(
        self,
        resources: dict[str, list[dict[str, Any]]],
        index: dict[str, dict[str, Any]] | None = None,
    ) -> None:
        self.resources = resources
        self.index = index or {}
        self.calls: list[tuple[str, dict[str, Any]]] = []

    async def get_resource(
        self, resource_type: str, params: dict[str, Any] | None = None
    ) -> dict[str, Any]:
        params = dict(params or {})
        self.calls.append((resource_type, params))
        items = self.resources.get(resource_type, [])
        limit = int(params["page[limit]"]) if "page[limit]" in params else (len(items) or 50)
        offset = int(params["page[offset]"]) if "page[offset]" in params else 0
        page = items[offset : offset + limit]
        included: list[dict[str, Any]] = []
        include = str(params.get("include") or "")
        if include:
            for item in page:
                relationships = item.get("relationships") or {}
                for rel_name in include.split(","):
                    rel = relationships.get(rel_name) or {}
                    data = rel.get("data")
                    rows = data if isinstance(data, list) else [data]
                    for row in rows:
                        if isinstance(row, dict) and row.get("id") in self.index:
                            included.append(self.index[str(row["id"])])
        return {"data": page, "included": included}


class _Drush:
    """run(*parts) answered by a handler receiving the same parts."""

    def __init__(self, handler: Any) -> None:
        self.handler = handler
        self.calls: list[tuple[Any, ...]] = []

    async def run(self, *parts: Any, **kwargs: Any) -> dict[str, Any]:
        self.calls.append(parts)
        return self.handler(*parts)

# Config collections carry UUID ids; the machine name lives in the
# drupal_internal__* field (type for node types, vid for vocabularies,
# id for media types).
_NODE_TYPES = [
    {"type": "node_type--node_type", "id": "uuid-nt-article",
     "attributes": {"drupal_internal__type": "article"}},
    {"type": "node_type--node_type", "id": "uuid-nt-page",
     "attributes": {"drupal_internal__type": "page"}},
]
_MEDIA_TYPES = [
    {"type": "media_type--media_type", "id": "uuid-mt-image",
     "attributes": {"drupal_internal__id": "image"}},
]
_VOCABS = [
    {"type": "taxonomy_vocabulary--taxonomy_vocabulary", "id": "uuid-tv-tags",
     "attributes": {"drupal_internal__vid": "tags"}},
]


def _node(
    nid: int,
    *,
    bundle: str = "article",
    title: str = "Title",
    status: bool = True,
    author: str | None = None,
    alias: str | None = None,
    extra: dict[str, Any] | None = None,
) -> dict[str, Any]:
    return {
        "type": f"node--{bundle}",
        "id": f"uuid-{nid}",
        "attributes": {
            "drupal_internal__nid": nid,
            "title": title,
            "status": status,
            "created": "2026-01-01T00:00:00+00:00",
            "changed": "2026-01-02T00:00:00+00:00",
            "langcode": "en",
            "path": {"alias": alias, "pid": None, "langcode": "en"},
            **(extra or {}),
        },
        "relationships": {
            "uid": {"data": ({"type": "user--user", "id": author} if author else None)}
        },
    }


def _read_csv(path: Path) -> tuple[list[str], list[list[str]]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv_module.reader(handle))
    return rows[0], rows[1:]


def _point_root(monkeypatch: Any, root: Path) -> None:
    monkeypatch.setattr(exports, "get_settings", lambda: _Settings(str(root)))


def test_export_nodes_writes_csv_with_author_and_path(monkeypatch, tmp_path):
    root = tmp_path / "site"
    root.mkdir()
    _point_root(monkeypatch, root)
    author_id = "author-uuid"
    api = _JSONAPI(
        {
            "node_type": _NODE_TYPES,
            "node--article": [
                _node(1, title="Published with alias", author=author_id, alias="/about"),
                _node(2, title="Unpublished", status=False),
                _node(3, title="Caf\u00e9 cr\u00e8me\nsecond line"),
            ],
            "node--page": [_node(10, bundle="page", title="A page")],
        },
        index={author_id: {"type": "user--user", "id": author_id, "attributes": {"drupal_internal__uid": 7, "name": "alice"}}},
    )
    fns = _tool_functions(exports, jsonapi_client=api)

    result = asyncio.run(fns["export_nodes"]())

    path = Path(result["file_path"])
    assert path.parent == root / "exports"
    assert path.name.startswith("nodes_all_") and path.suffix == ".csv"
    assert result["rows"] == 3  # the unpublished node was skipped
    assert result["columns"] == 11
    assert result["truncated"] is False
    assert 0 < result["file_size_kb"]
    assert len(result["preview"]) == 3

    columns, rows = _read_csv(path)
    assert columns == exports._NODE_COLUMNS
    assert rows[0][0] == "1" and rows[0][2] == "article"
    assert rows[0][8] == "7" and rows[0][9] == "alice"  # author from include
    assert rows[0][10] == "/about"
    assert rows[1][0] == "3"  # the unpublished node (2) never made it to the file
    assert rows[1][4] == "published"
    assert rows[1][8] == "" and rows[1][9] == ""  # no author relationship
    assert rows[1][10] == "/node/3"  # no alias: the internal path
    assert rows[1][3] == "Caf\u00e9 cr\u00e8me\nsecond line"  # newline intact
    assert rows[2][0] == "10" and rows[2][2] == "page"
    assert rows[2][10] == "/node/10"
    # bundle discovery walked node_type, then one endpoint per bundle
    walked = [name for name, _ in api.calls]
    assert walked[0] == "node_type"
    assert "node--article" in walked and "node--page" in walked
    assert "node" not in walked, "bare merged collection must never be paged"


def test_export_nodes_include_unpublished_and_fields_column(monkeypatch, tmp_path):
    root = tmp_path / "site"
    root.mkdir()
    _point_root(monkeypatch, root)
    long_body = "x" * 6000
    api = _JSONAPI(
        {
            "node_type": _NODE_TYPES,
            "node--article": [
                _node(1, status=False, extra={"body": {"value": long_body}, "promote": True}),
            ],
        }
    )
    fns = _tool_functions(exports, jsonapi_client=api)

    result = asyncio.run(
        fns["export_nodes"](include_unpublished=True, include_fields=True)
    )

    columns, rows = _read_csv(Path(result["file_path"]))
    assert columns == exports._NODE_COLUMNS + ["fields"]
    assert len(rows) == 1
    fields = json.loads(rows[0][11])
    assert fields["promote"] is True
    assert fields["body"]["value"].startswith("x" * 5000)
    assert "[truncated]" in fields["body"]["value"]
    assert len(fields["body"]["value"]) < 5100
    assert result["filters"] == {
        "content_type": None,
        "include_unpublished": True,
        "include_fields": True,
        "limit": 0,
    }


def test_export_nodes_limit_trims_and_flags_truncation(monkeypatch, tmp_path):
    root = tmp_path / "site"
    root.mkdir()
    _point_root(monkeypatch, root)
    api = _JSONAPI(
        {
            "node_type": _NODE_TYPES,
            "node--article": [_node(n) for n in range(1, 6)],
        }
    )
    fns = _tool_functions(exports, jsonapi_client=api)

    result = asyncio.run(fns["export_nodes"](limit=2))

    assert result["rows"] == 2
    assert result["truncated"] is True
    assert "incomplete" in result["note"]
    _, rows = _read_csv(Path(result["file_path"]))
    assert [row[0] for row in rows] == ["1", "2"]


def test_export_nodes_pagination_walks_every_page(monkeypatch, tmp_path):
    root = tmp_path / "site"
    root.mkdir()
    _point_root(monkeypatch, root)
    api = _JSONAPI(
        {
            "node_type": _NODE_TYPES,
            "node--article": [_node(n) for n in range(1, 451)],
        }
    )
    fns = _tool_functions(exports, jsonapi_client=api)

    result = asyncio.run(fns["export_nodes"]())

    assert result["rows"] == 450
    offsets = [
        params["page[offset]"]
        for name, params in api.calls
        if name == "node--article"
    ]
    assert offsets == ["0", "200", "400"]  # 50-row last page ends the walk
    assert result["truncated"] is False


def test_export_nodes_rejects_bad_inputs_and_unsafe_paths(monkeypatch, tmp_path):
    _point_root(monkeypatch, tmp_path / "site")
    api = _JSONAPI({"node_type": _NODE_TYPES, "node--article": [_node(1)]})
    fns = _tool_functions(exports, jsonapi_client=api)

    assert "error" in asyncio.run(fns["export_nodes"](content_type="article; DROP TABLE x"))
    assert "error" in asyncio.run(fns["export_nodes"](format="pdf"))
    assert "error" in asyncio.run(fns["export_nodes"](limit=-1))
    assert "error" in asyncio.run(fns["export_nodes"](limit=True))
    assert "error" in asyncio.run(fns["export_nodes"](limit=100001))

    outside = asyncio.run(fns["export_nodes"](output_path="/etc/drupal-export.csv"))
    assert "must live inside" in outside.get("error", ""), outside
    bad_ext = asyncio.run(fns["export_nodes"](output_path=str(tmp_path / "out.txt")))
    assert "must end in .csv or .xlsx" in bad_ext.get("error", ""), bad_ext
    mismatch = asyncio.run(fns["export_nodes"](output_path=str(tmp_path / "out.xlsx")))
    assert "format is 'csv'" in mismatch.get("error", ""), mismatch


def test_export_nodes_no_rows_writes_no_file(monkeypatch, tmp_path):
    root = tmp_path / "site"
    root.mkdir()
    _point_root(monkeypatch, root)
    api = _JSONAPI({"node_type": _NODE_TYPES, "node--article": []})
    fns = _tool_functions(exports, jsonapi_client=api)

    result = asyncio.run(fns["export_nodes"]())

    assert result == {
        "rows": 0,
        "note": "No nodes matched; no file was written.",
        "filters": {
            "content_type": None,
            "include_unpublished": False,
            "include_fields": False,
            "limit": 0,
        },
        "skipped_without_id": 0,
    }
    exports_dir = root / "exports"
    assert not exports_dir.exists() or not list(exports_dir.glob("*"))


def test_export_nodes_unknown_bundle_writes_no_file(monkeypatch, tmp_path):
    root = tmp_path / "site"
    root.mkdir()
    _point_root(monkeypatch, root)
    api = _JSONAPI({"node_type": _NODE_TYPES, "node--event": []})
    fns = _tool_functions(exports, jsonapi_client=api)

    result = asyncio.run(fns["export_nodes"](content_type="event"))

    assert result["rows"] == 0
    assert "no file" in result["note"]


def test_export_users_xlsx_roundtrip_with_roles(monkeypatch, tmp_path):
    root = tmp_path / "site"
    root.mkdir()
    _point_root(monkeypatch, root)
    role_id = "role-uuid"
    api = _JSONAPI(
        {
            "user--user": [
                {
                    "type": "user--user",
                    "id": "u1",
                    "attributes": {
                        "drupal_internal__uid": 1,
                        "name": "admin",
                        "mail": "admin@example.com",
                        "status": True,
                        "created": "2026-01-01T00:00:00+00:00",
                        "access": "2026-02-01T00:00:00+00:00",
                        "login": "2026-02-01T00:00:00+00:00",
                        "langcode": "en",
                    },
                    "relationships": {"roles": {"data": [{"type": "user_role--user_role", "id": role_id}]}},
                },
                {
                    "type": "user--user",
                    "id": "u2",
                    "attributes": {
                        "drupal_internal__uid": 2,
                        "name": "blocked_bob",
                        "mail": "bob@example.com",
                        "status": False,
                        "created": "2026-01-01T00:00:00+00:00",
                        "access": None,
                        "login": None,
                        "langcode": "en",
                    },
                    "relationships": {"roles": {"data": []}},
                },
            ]
        },
        index={role_id: {"type": "user_role--user_role", "id": role_id, "attributes": {"label": "Administrator"}}},
    )
    fns = _tool_functions(exports, jsonapi_client=api)

    out = tmp_path / "users.xlsx"
    result = asyncio.run(fns["export_users"](output_path=str(out), format="xlsx"))

    assert result["file_path"] == str(out.resolve())
    assert result["rows"] == 2

    from openpyxl import load_workbook

    sheet = load_workbook(out)["export"]
    rows = list(sheet.iter_rows(values_only=True))
    assert list(rows[0]) == exports._USER_COLUMNS
    assert rows[1][1] == "admin" and rows[1][4] == "Administrator"
    # openpyxl reads a cell written as "" back as None
    assert rows[2][3] == "blocked" and not rows[2][4]
    # include=roles was actually requested
    assert api.calls[0][1].get("include") == "roles"

def test_export_users_skips_entries_without_an_internal_id(monkeypatch, tmp_path):
    """Drupal's uid-0 anonymous user serves only a display_name; exporting
    it would put a row of nulls in the account inventory."""
    _point_root(monkeypatch, tmp_path)
    api = _JSONAPI(
        {
            "user--user": [
                {"type": "user--user", "id": "anon-uuid", "attributes": {"display_name": "Anonymous"}},
                {
                    "type": "user--user",
                    "id": "u1",
                    "attributes": {
                        "drupal_internal__uid": 1,
                        "name": "admin",
                        "mail": "admin@example.com",
                        "status": True,
                        "created": "2026-01-01T00:00:00+00:00",
                        "access": None,
                        "login": None,
                        "langcode": "en",
                    },
                    "relationships": {"roles": {"data": []}},
                },
            ]
        }
    )
    fns = _tool_functions(exports, jsonapi_client=api)

    result = asyncio.run(fns["export_users"]())

    assert result["rows"] == 1
    assert result["skipped_without_id"] == 1
    _, rows = _read_csv(Path(result["file_path"]))
    assert len(rows) == 1
    assert rows[0][1] == "admin"


def test_export_media_lists_file_details(monkeypatch, tmp_path):
    root = tmp_path / "site"
    root.mkdir()
    _point_root(monkeypatch, root)
    file_id = "file-uuid"
    owner_id = "owner-uuid"
    api = _JSONAPI(
        {
            "media_type": _MEDIA_TYPES,
            "media--image": [
                {
                    "type": "media--image",
                    "id": "m1",
                    "attributes": {
                        "drupal_internal__mid": 4,
                        "name": "Hero",
                        "status": True,
                        "created": "2026-01-01T00:00:00+00:00",
                        "changed": "2026-01-02T00:00:00+00:00",
                    },
                    "relationships": {
                        "uid": {"data": {"type": "user--user", "id": owner_id}},
                        "field_media_file": {"data": {"type": "file--file", "id": file_id}},
                    },
                }
            ],
        },
        index={
            file_id: {
                "type": "file--file",
                "id": file_id,
                "attributes": {"filename": "hero.jpg", "uri": "public://hero.jpg"},
            },
            owner_id: {
                "type": "user--user",
                "id": owner_id,
                "attributes": {"drupal_internal__uid": 7, "name": "alice"},
            },
        },
    )
    fns = _tool_functions(exports, jsonapi_client=api)

    result = asyncio.run(fns["export_media"]())

    columns, rows = _read_csv(Path(result["file_path"]))
    assert columns == exports._MEDIA_COLUMNS
    assert rows[0] == [
        "4", "image", "Hero", "active",
        "2026-01-01T00:00:00+00:00", "2026-01-02T00:00:00+00:00",
        "alice", "hero.jpg", "public://hero.jpg",
    ]


def test_export_media_empty_site_says_so(monkeypatch, tmp_path):
    _point_root(monkeypatch, tmp_path)
    api = _JSONAPI({"media_type": _MEDIA_TYPES, "media--image": []})
    fns = _tool_functions(exports, jsonapi_client=api)

    result = asyncio.run(fns["export_media"]())

    assert result["rows"] == 0
    assert "No media" in result["note"]


def test_export_taxonomy_counts_usage_with_one_query_per_field(monkeypatch, tmp_path):
    root = tmp_path / "site"
    root.mkdir()
    _point_root(monkeypatch, root)
    api = _JSONAPI(
        {
            "taxonomy_vocabulary": _VOCABS,
            "taxonomy_term--tags": [
                {
                    "type": "taxonomy_term--tags",
                    "id": "t1",
                    "attributes": {
                        "drupal_internal__tid": 1,
                        "name": "Used term",
                        "description": None,
                        "weight": 0,
                        "status": True,
                    },
                },
                {
                    "type": "taxonomy_term--tags",
                    "id": "t2",
                    "attributes": {
                        "drupal_internal__tid": 2,
                        "name": "Orphan term",
                        "description": None,
                        "weight": 1,
                        "status": True,
                    },
                },
            ],
            "field_storage_config": [
                {
                    "type": "field_storage_config--field_storage_config",
                    "id": "node.field_tags",
                    "attributes": {
                        "entity_type": "node",
                        "field_name": "field_tags",
                        "settings": {"target_type": "taxonomy_term"},
                    },
                },
                {
                    "type": "field_storage_config--field_storage_config",
                    "id": "node.field_image",
                    "attributes": {
                        "entity_type": "node",
                        "field_name": "field_image",
                        "settings": {"target_type": "file"},
                    },
                },
                {
                    "type": "field_storage_config--field_storage_config",
                    "id": "bad.field",
                    "attributes": {
                        "entity_type": "node; DROP TABLE users",
                        "field_name": "field_tags",
                        "settings": {"target_type": "taxonomy_term"},
                    },
                },
            ],
        }
    )

    def handler(*parts: Any) -> dict[str, Any]:
        command, sql = parts[0], str(parts[1])
        assert command == "sql:query"
        if sql == "SELECT target_id, COUNT(*) AS uses FROM node__field_tags GROUP BY target_id":
            return {"returncode": 0, "data": {"output": "1\t3\n2\t0\n"}}
        return {"returncode": 1, "stderr": "query failed", "data": {}}

    drush = _Drush(handler)
    fns = _tool_functions(exports, jsonapi_client=api, drush_client=drush)

    result = asyncio.run(fns["export_taxonomy"]())

    columns, rows = _read_csv(Path(result["file_path"]))
    assert columns == exports._TERM_COLUMNS
    assert rows[0][0] == "1" and rows[0][6] == "3"
    assert rows[0][7] == "node.field_tags"
    # a counting gap was reported, so the verdict is withheld for every row
    assert rows[0][8] == ""
    assert rows[1][6] == "0" and rows[1][8] == ""
    # exactly one COUNT query, only for the taxonomy-targeting storage
    queries = [parts[1] for parts in drush.calls if parts[0] == "sql:query"]
    assert len(queries) == 1
    assert "node; DROP" not in str(drush.calls)
    # the invalid entity_type was reported rather than silently dropped
    assert "node; DROP TABLE users" in result["usage_note"]
    assert "field_name,entity_type,settings" in str(
        [params for name, params in api.calls if name == "field_storage_config"]
    )


def test_export_taxonomy_verdict_when_counting_is_complete(monkeypatch, tmp_path):
    _point_root(monkeypatch, tmp_path)
    api = _JSONAPI(
        {
            "taxonomy_vocabulary": _VOCABS,
            "taxonomy_term--tags": [
                {
                    "type": "taxonomy_term--tags",
                    "id": "t1",
                    "attributes": {"drupal_internal__tid": 1, "name": "Used", "description": None, "weight": 0, "status": True},
                },
                {
                    "type": "taxonomy_term--tags",
                    "id": "t2",
                    "attributes": {"drupal_internal__tid": 2, "name": "Free", "description": None, "weight": 1, "status": True},
                },
            ],
            "field_storage_config": [
                {
                    "type": "field_storage_config--field_storage_config",
                    "id": "node.field_tags",
                    "attributes": {
                        "entity_type": "node",
                        "field_name": "field_tags",
                        "settings": {"target_type": "taxonomy_term"},
                    },
                }
            ],
        }
    )
    drush = _Drush(
        lambda *parts: {"returncode": 0, "data": {"output": "1\t2\n"}}
    )
    fns = _tool_functions(exports, jsonapi_client=api, drush_client=drush)

    result = asyncio.run(fns["export_taxonomy"](vocabulary="tags"))

    assert "usage_note" not in result
    _, rows = _read_csv(Path(result["file_path"]))
    assert rows[0][6] == "2" and rows[0][8] == "no"
    assert rows[1][6] == "0" and rows[1][8] == "yes"
    # usage was computed first, then exactly the requested bundle's terms
    assert [name for name, _ in api.calls] == ["field_storage_config", "taxonomy_term--tags"]


def test_export_taxonomy_without_drush_degrades_with_a_note(monkeypatch, tmp_path):
    _point_root(monkeypatch, tmp_path)
    api = _JSONAPI(
        {
            "taxonomy_vocabulary": _VOCABS,
            "taxonomy_term--tags": [
                {
                    "type": "taxonomy_term--tags",
                    "id": "t1",
                    "attributes": {
                        "drupal_internal__tid": 1,
                        "name": "Term",
                        "description": None,
                        "weight": 0,
                        "status": True,
                    },
                }
            ],
        }
    )
    fns = _tool_functions(exports, jsonapi_client=api)  # no drush_client

    result = asyncio.run(fns["export_taxonomy"]())

    assert "Drush" in result["usage_note"]
    _, rows = _read_csv(Path(result["file_path"]))
    assert rows[0][6:] == ["", "", ""]


def test_export_taxonomy_usage_can_be_skipped(monkeypatch, tmp_path):
    _point_root(monkeypatch, tmp_path)
    api = _JSONAPI(
        {
            "taxonomy_vocabulary": _VOCABS,
            "taxonomy_term--tags": [
                {
                    "type": "taxonomy_term--tags",
                    "id": "t1",
                    "attributes": {
                        "drupal_internal__tid": 1,
                        "name": "Term",
                        "description": None,
                        "weight": 0,
                        "status": True,
                    },
                }
            ],
        }
    )
    drush = _Drush(lambda *parts: {"returncode": 0, "data": {"output": ""}})
    fns = _tool_functions(exports, jsonapi_client=api, drush_client=drush)

    result = asyncio.run(fns["export_taxonomy"](include_usage=False))

    assert "usage_note" not in result
    assert drush.calls == []
    assert api.calls[-1][0] == "taxonomy_term--tags"


def test_export_taxonomy_vocabulary_validation(monkeypatch, tmp_path):
    _point_root(monkeypatch, tmp_path)
    api = _JSONAPI({})
    fns = _tool_functions(exports, jsonapi_client=api)

    result = asyncio.run(fns["export_taxonomy"](vocabulary="tags; DROP TABLE x"))

    assert "error" in result and "vocabulary" in result["error"]
    assert api.calls == []


def test_export_defaults_to_temp_when_drupal_root_is_unknown(monkeypatch, tmp_path):
    def boom() -> Any:
        raise RuntimeError("settings unavailable")

    monkeypatch.setattr(exports, "get_settings", boom)
    api = _JSONAPI({"node_type": _NODE_TYPES, "node--article": [_node(1)]})
    fns = _tool_functions(exports, jsonapi_client=api)

    result = asyncio.run(fns["export_nodes"]())

    path = Path(result["file_path"])
    expected_parent = Path(tempfile.gettempdir()).resolve() / "drupal-mcp-exports"
    assert path.parent == expected_parent
    assert path.is_relative_to(Path(tempfile.gettempdir()).resolve())


def test_all_four_tools_error_without_a_jsonapi_client():
    fns = _tool_functions(exports)  # no clients at all

    for name in ("export_nodes", "export_users", "export_media", "export_taxonomy"):
        result = asyncio.run(fns[name]())
        assert "error" in result, name
        assert "JSON:API" in result["error"], name
